"""Generate the FICTIONAL mapping fixtures and their ground-truth mappings.

All companies, people and addresses are invented; e-mail domains are .example/.invalid;
the only IBAN is the documented test IBAN DE89370400440532013000.
Run from the repository root:  python mapping/fixtures/make_fixtures.py
The generated files are committed; re-running reproduces them byte for byte.
"""
from __future__ import annotations

import csv
import io
import json
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

HERE = Path(__file__).resolve().parent

BUYERS = {
    "baecker": dict(name="Bäckerei Sonnenschein GmbH", street="Lindenstraße 12", postcode="10115", city="Berlin", country="DE",
                    email="buchhaltung@baeckerei-sonnenschein.example"),
    "mueller": dict(name="Müller & Söhne Haustechnik KG", street="Am Markt 3", postcode="01067", city="Dresden", country="DE",
                    email="rechnung@mueller-haustechnik.example"),
    "stadt": dict(name="Stadtverwaltung Musterstadt", street="Rathausplatz 1", postcode="99999", city="Musterstadt", country="DE",
                  email="e-rechnung@musterstadt.invalid", leitweg="04011000-1234512345-06"),
    "cafe": dict(name="Café Zeitlos UG (haftungsbeschränkt)", street="Hauptstraße 45a", postcode="80331", city="München", country="DE",
                 email="info@cafe-zeitlos.example"),
    "nordlicht": dict(name="Nordlicht Design GmbH", street="Hafenweg 7", postcode="20457", city="Hamburg", country="DE",
                      email="office@nordlicht-design.example", vat_id="DE999999999"),
    "alpen": dict(name="Alpenblick Hotel GmbH", street="Seestraße 2", postcode="6020", city="Innsbruck", country="AT",
                  email="einkauf@alpenblick-hotel.invalid", vat_id="ATU99999999"),
}

C2 = Decimal("0.01")


def q(d: Decimal) -> Decimal:
    return d.quantize(C2, rounding=ROUND_HALF_UP)


def de(d: Decimal, thousands: bool = True) -> str:
    s = f"{q(d):,.2f}" if thousands else f"{q(d):.2f}"
    return s.replace(",", "_").replace(".", ",").replace("_", ".")


def en(d: Decimal, thousands: bool = False) -> str:
    return f"{q(d):,.2f}" if thousands else f"{q(d):.2f}"


def csv_text(rows: list[list[str]], delimiter: str, quoting=csv.QUOTE_MINIMAL) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=delimiter, quoting=quoting, lineterminator="\r\n")
    w.writerows(rows)
    return buf.getvalue()


def col(target: str, column: str, transform: list[str] | None = None) -> dict:
    return {"target": target, "source_column": column, "transform": transform or []}


def const(target: str, value: str) -> dict:
    return {"target": target, "constant": value}


def write_expected(name: str, mode: str, number_column: str | None, fields: list[dict]) -> None:
    data = {"fixture": name, "grouping": {"mode": mode, "invoice_number_column": number_column}, "fields": fields}
    (HERE / (Path(name).stem + ".expected_mapping.json")).write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def buyer_fields(name, street, postcode, city, email, country=None) -> list[dict]:
    out = [col("/buyer/name", name), col("/buyer/street", street), col("/buyer/postcode", postcode), col("/buyer/city", city),
           col("/buyer/electronic_address", email), const("/buyer/electronic_address_scheme", "EM")]
    out.append(col("/buyer/country", country) if country else const("/buyer/country", "DE"))
    return out


# ------------------------------------------------------------ 1 lexware-like: ';', CP1252, German decimals, row per invoice
def lexware():
    header = ["Rechnungsnr.", "Rechnungsdatum", "Kundennr.", "Kunde", "Straße", "PLZ", "Ort", "Land", "E-Mail", "Leistungsdatum",
              "Bezeichnung", "Menge", "Einheit", "Einzelpreis", "Netto", "MwSt-Satz", "MwSt-Betrag", "Brutto"]
    data = [
        ("RE-2026-0101", "02.10.2026", "10001", "baecker", "30.09.2026", "Wartung Kühlanlage", "3", "Std", "85.00", 19),
        ("RE-2026-0102", "05.10.2026", "10002", "mueller", "01.10.2026", "Montage Lüftungsanlage pauschal", "1", "Pauschal", "1250.00", 19),
        ("RE-2026-0103", "09.10.2026", "10003", "cafe", "08.10.2026", "Ersatzteil Dichtungssatz", "4", "Stk", "12.50", 19),
        ("RE-2026-0104", "15.10.2026", "10001", "baecker", "14.10.2026", "Notdienst Wochenende", "2,5", "Std", "110.00", 19),
        ("RE-2026-0105", "20.10.2026", "10005", "nordlicht", "19.10.2026", "Inspektion Klimagerät", "1", "Stk", "189.00", 19),
    ]
    rows = [header]
    for num, d, kd, b, ld, text, qty, unit, price, rate in data:
        B = BUYERS[b]
        net = q(Decimal(qty.replace(",", ".")) * Decimal(price))
        vat = q(net * rate / 100)
        rows.append([num, d, kd, B["name"], B["street"], B["postcode"], B["city"], "Deutschland", B["email"], ld, text, qty, unit,
                     de(Decimal(price)), de(net), f"{rate} %", de(vat), de(net + vat)])
    (HERE / "lexware_like.csv").write_bytes(csv_text(rows, ";").encode("cp1252"))
    write_expected("lexware_like.csv", "row_per_invoice", "Rechnungsnr.", [
        col("/number", "Rechnungsnr."), col("/issue_date", "Rechnungsdatum", ["date_dmy"]),
        col("/delivery_date", "Leistungsdatum", ["date_dmy"]),
        *buyer_fields("Kunde", "Straße", "PLZ", "Ort", "E-Mail", country="Land"),
        col("/lines/*/description", "Bezeichnung"), col("/lines/*/quantity", "Menge", ["de_decimal"]),
        col("/lines/*/unit", "Einheit", ["unit_map"]), col("/lines/*/net_price", "Einzelpreis", ["de_decimal"]),
        col("/lines/*/vat_rate", "MwSt-Satz", ["de_decimal"]), col("/lines/*/vat_category", "MwSt-Satz", ["rate_to_category"]),
        col("/source/expected_totals/net", "Netto", ["de_decimal"]), col("/source/expected_totals/vat", "MwSt-Betrag", ["de_decimal"]),
        col("/source/expected_totals/gross", "Brutto", ["de_decimal"]),
    ])


# ------------------------------------------------------------ 2 sevdesk-like: ',', UTF-8 BOM, English decimals, ISO dates
def sevdesk():
    header = ["Invoice number", "Invoice date", "Due date", "Customer", "Street", "ZIP", "City", "Country", "Email", "VAT ID",
              "Description", "Quantity", "Unit", "Unit price net", "VAT %", "Total net", "Total gross"]
    data = [
        ("INV-1001", "2026-10-05", "2026-10-19", "nordlicht", "Website-Wartung Oktober", "1", "pcs", "1200.00", 19),
        ("INV-1002", "2026-10-07", "2026-10-21", "alpen", "Fotoshooting Hotelzimmer", "6", "hours", "95.00", 19),
        ("INV-1003", "2026-10-12", "2026-10-26", "cafe", "Speisekarten-Layout", "1", "pcs", "340.00", 19),
        ("INV-1004", "2026-10-28", "2026-11-11", "mueller", "Logo-Überarbeitung", "8", "hours", "75.00", 19),
    ]
    rows = [header]
    for num, d, due, b, text, qty, unit, price, rate in data:
        B = BUYERS[b]
        net = q(Decimal(qty) * Decimal(price))
        rows.append([num, d, due, B["name"], B["street"], B["postcode"], B["city"], B["country"], B["email"], B.get("vat_id", ""),
                     text, qty, unit, en(Decimal(price), thousands=True), str(rate), en(net), en(q(net * (100 + rate) / 100))])
    (HERE / "sevdesk_like.csv").write_bytes(b"\xef\xbb\xbf" + csv_text(rows, ",", csv.QUOTE_ALL).encode("utf-8"))
    write_expected("sevdesk_like.csv", "row_per_invoice", "Invoice number", [
        col("/number", "Invoice number"), col("/issue_date", "Invoice date", ["date_iso"]),
        col("/payment/due_date", "Due date", ["date_iso"]),
        *buyer_fields("Customer", "Street", "ZIP", "City", "Email", country="Country"),
        col("/buyer/vat_id", "VAT ID"),
        col("/lines/*/description", "Description"), col("/lines/*/quantity", "Quantity", ["en_decimal"]),
        col("/lines/*/unit", "Unit", ["unit_map"]), col("/lines/*/net_price", "Unit price net", ["en_decimal"]),
        col("/lines/*/vat_rate", "VAT %", ["en_decimal"]), col("/lines/*/vat_category", "VAT %", ["rate_to_category"]),
        col("/source/expected_totals/net", "Total net", ["en_decimal"]), col("/source/expected_totals/gross", "Total gross", ["en_decimal"]),
    ])


# ------------------------------------------------------------ 3 hand-made Excel: real dates/floats/percent, PLZ as number
def excel():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rechnungen 2026"
    header = ["Re-Nr", "Datum", "Kunde", "Straße + Nr.", "PLZ", "Ort", "Mail", "Leistung", "Stunden", "Stundensatz", "MwSt", "Gesamt brutto"]
    ws.append(header)
    data = [
        (26041, datetime(2026, 10, 1), "mueller", "Beratung Heizungstausch", 7.5, 92.5),
        (26042, datetime(2026, 10, 8), "stadt", "Gutachten Dachstuhl", 12, 105),
        (26043, datetime(2026, 10, 16), "baecker", "Energieberatung vor Ort", 3.25, 92.5),
    ]
    for num, d, b, text, hours, rate in data:
        B = BUYERS[b]
        gross = float(q(Decimal(str(hours)) * Decimal(str(rate)) * Decimal("1.19")))
        pc = int(B["postcode"]) if B["postcode"].isdigit() else B["postcode"]  # Excel turns 01067 into 1067
        ws.append([num, d, B["name"], B["street"], pc, B["city"], B["email"], text, hours, rate, 0.19, gross])
        r = ws.max_row
        ws.cell(r, 2).number_format = "DD.MM.YYYY"
        ws.cell(r, 11).number_format = "0%"
        ws.cell(r, 12).number_format = "#,##0.00 €"
    ws.append([])  # trailing empty row
    buf = io.BytesIO()
    wb.save(buf)
    _normalise_zip(buf.getvalue(), HERE / "excel_handmade.xlsx")
    write_expected("excel_handmade.xlsx", "row_per_invoice", "Re-Nr", [
        col("/number", "Re-Nr"), col("/issue_date", "Datum", ["date_dmy"]),
        *buyer_fields("Kunde", "Straße + Nr.", "PLZ", "Ort", "Mail"),
        col("/lines/*/description", "Leistung"), col("/lines/*/quantity", "Stunden", ["en_decimal"]),
        const("/lines/*/unit", "HUR"), col("/lines/*/net_price", "Stundensatz", ["en_decimal"]),
        col("/lines/*/vat_rate", "MwSt", ["en_decimal"]), col("/lines/*/vat_category", "MwSt", ["rate_to_category"]),
        col("/source/expected_totals/gross", "Gesamt brutto", ["en_decimal"]),
    ])


def _normalise_zip(data: bytes, path: Path) -> None:
    """Rewrite the xlsx with fixed timestamps so regeneration is reproducible."""
    import zipfile
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            content = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                import re
                content = re.sub(rb"<dcterms:(created|modified)([^>]*)>[^<]*<", rb"<dcterms:\1\2>2026-10-31T00:00:00Z<", content)
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 10, 31, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(zi, content)
    path.write_bytes(out.getvalue())


# ------------------------------------------------------------ 4 Access-style export: title rows, tab, abbreviations, date+time
def access():
    header = ["RechNr", "RechDatum", "KdName", "KdStrasse", "KdPLZ", "KdOrt", "KdEmail", "ArtBez", "Anz", "EPreis", "UStProz", "LeitwegID"]
    data = [
        ("A-7781", "06.10.2026", "stadt", "Reinigung Bürgersaal", "1", "480.00", 19),
        ("A-7782", "13.10.2026", "cafe", "Fensterreinigung", "2", "65.00", 19),
        ("A-7783", "20.10.2026", "stadt", "Grundreinigung Turnhalle", "1", "1150.00", 19),
        ("A-7784", "27.10.2026", "baecker", "Unterhaltsreinigung Oktober", "4", "120.00", 19),
    ]
    rows = [["Rechnungsausgangsbuch – Export aus Auftragsverwaltung"], ["Erstellt: 31.10.2026 08:15"], [], header]
    for num, d, b, text, qty, price, rate in data:
        B = BUYERS[b]
        rows.append([num, d + " 00:00:00", B["name"], B["street"], B["postcode"], B["city"], B["email"], text, qty,
                     de(Decimal(price), thousands=False), str(rate), B.get("leitweg", "")])
    rows.append([])
    rows.append([])
    (HERE / "access_export.csv").write_bytes(csv_text(rows, "\t").encode("utf-8"))
    write_expected("access_export.csv", "row_per_invoice", "RechNr", [
        col("/number", "RechNr"), col("/issue_date", "RechDatum", ["date_dmy"]),
        *buyer_fields("KdName", "KdStrasse", "KdPLZ", "KdOrt", "KdEmail"),
        col("/buyer_reference", "LeitwegID"),
        col("/lines/*/description", "ArtBez"), col("/lines/*/quantity", "Anz", ["de_decimal"]),
        const("/lines/*/unit", "C62"), col("/lines/*/net_price", "EPreis", ["de_decimal"]),
        col("/lines/*/vat_rate", "UStProz", ["de_decimal"]), col("/lines/*/vat_category", "UStProz", ["rate_to_category"]),
    ])


# ------------------------------------------------------------ 5 one row per line, repeated invoice numbers
def per_line():
    header = ["Belegnr", "Belegdatum", "Pos", "Kunde", "Anschrift", "PLZ", "Ort", "E-Mail", "Bestellnummer", "Käuferreferenz",
              "Artikel", "Menge", "ME", "Einzelpreis netto", "USt-Satz", "Positionsbetrag netto"]
    invoices = [
        ("2026-301", "03.10.2026", "stadt", "B-55120", "04011000-1234512345-06", [
            ("Druck Broschüre A5, 500 Stück", "500", "Stk", "0.84"),
            ("Satz und Layout", "6", "Std", "68.00"),
            ("Versandpauschale", "1", "Pauschal", "14.90"),
        ]),
        ("2026-302", "10.10.2026", "cafe", "", "", [
            ("Visitenkarten 250 Stück", "250", "Stk", "0.12"),
            ("Flyer DIN lang 1.000 Stück", "1000", "Stk", "0.09"),
        ]),
        ("2026-303", "24.10.2026", "nordlicht", "PO-8812", "NL-2026-17", [
            ("Plakate A1 (Großformat)", "20", "Stk", "18.50"),
        ]),
    ]
    rows = [header]
    for num, d, b, order, ref, lines in invoices:
        B = BUYERS[b]
        for i, (text, qty, unit, price) in enumerate(lines, 1):
            rows.append([num, d, str(i), B["name"], B["street"], B["postcode"], B["city"], B["email"], order, ref, text,
                         de(Decimal(qty), thousands=True).removesuffix(",00"), unit, de(Decimal(price)), "19,00",
                         de(Decimal(qty) * Decimal(price))])
    (HERE / "row_per_line.csv").write_bytes(csv_text(rows, ";").encode("utf-8"))
    write_expected("row_per_line.csv", "row_per_line", "Belegnr", [
        col("/number", "Belegnr"), col("/issue_date", "Belegdatum", ["date_dmy"]),
        *buyer_fields("Kunde", "Anschrift", "PLZ", "Ort", "E-Mail"),
        col("/order_reference", "Bestellnummer"), col("/buyer_reference", "Käuferreferenz"),
        col("/lines/*/id", "Pos"), col("/lines/*/description", "Artikel"), col("/lines/*/quantity", "Menge", ["de_decimal"]),
        col("/lines/*/unit", "ME", ["unit_map"]), col("/lines/*/net_price", "Einzelpreis netto", ["de_decimal"]),
        col("/lines/*/vat_rate", "USt-Satz", ["de_decimal"]), col("/lines/*/vat_category", "USt-Satz", ["rate_to_category"]),
    ])


# ------------------------------------------------------------ 6 no VAT column (constant 19 % from the seller profile)
def no_vat():
    header = ["Belegnummer", "Datum", "Fällig am", "Firma", "Straße", "PLZ", "Ort", "E-Mail", "Leistung", "Anzahl", "Preis"]
    data = [
        ("2026/0451", "01.10.2026", "15.10.2026", "mueller", "Gerüstmiete Woche 40", "1", "640,00"),
        ("2026/0452", "08.10.2026", "22.10.2026", "baecker", "Container 7 m³ inkl. Entsorgung", "2", "385,00"),
        ("2026/0453", "22.10.2026", "05.11.2026", "cafe", "Transport Baumaterial", "3", "79,90"),
    ]
    rows = [header]
    for num, d, due, b, text, qty, price in data:
        B = BUYERS[b]
        rows.append([num, d, due, B["name"], B["street"], B["postcode"], B["city"], B["email"], text, qty, price])
    (HERE / "no_vat_column.csv").write_bytes(csv_text(rows, ";").encode("cp1252"))
    write_expected("no_vat_column.csv", "row_per_invoice", "Belegnummer", [
        col("/number", "Belegnummer"), col("/issue_date", "Datum", ["date_dmy"]), col("/payment/due_date", "Fällig am", ["date_dmy"]),
        *buyer_fields("Firma", "Straße", "PLZ", "Ort", "E-Mail"),
        col("/lines/*/description", "Leistung"), col("/lines/*/quantity", "Anzahl", ["de_decimal"]),
        const("/lines/*/unit", "C62"), col("/lines/*/net_price", "Preis", ["de_decimal"]),
        const("/lines/*/vat_rate", "19"), const("/lines/*/vat_category", "S"),
    ])


# ------------------------------------------------------------ 7 mixed 7/19 % lines with GROSS unit prices
def mixed_vat():
    header = ["Rechnungsnummer", "Rechnungsdatum", "Lieferdatum", "Kunde", "Straße", "PLZ", "Ort", "E-Mail", "Artikel", "Menge",
              "Einheit", "Einzelpreis brutto", "MwSt %", "Gesamt brutto"]
    invoices = [
        ("CAT-2026-88", "12.10.2026", "11.10.2026", "stadt", [
            ("Brötchen belegt", "40", "Stk", "2.14", 7),
            ("Kaffee (Kanne)", "6", "Stk", "8.56", 7),
            ("Miete Kaffeemaschine", "1", "Tag", "35.70", 19),
            ("Servicepersonal", "4", "Std", "41.65", 19),
        ]),
        ("CAT-2026-89", "19.10.2026", "18.10.2026", "nordlicht", [
            ("Kuchenplatte", "2", "Stk", "26.75", 7),
            ("Lieferung", "1", "Pauschal", "23.80", 19),
        ]),
    ]
    rows = [header]
    for num, d, ld, b, lines in invoices:
        B = BUYERS[b]
        for text, qty, unit, gross, rate in lines:
            rows.append([num, d, ld, B["name"], B["street"], B["postcode"], B["city"], B["email"], text, qty, unit,
                         de(Decimal(gross)), str(rate), de(Decimal(qty) * Decimal(gross))])
    (HERE / "mixed_vat.csv").write_bytes(csv_text(rows, ";").encode("utf-8"))
    write_expected("mixed_vat.csv", "row_per_line", "Rechnungsnummer", [
        col("/number", "Rechnungsnummer"), col("/issue_date", "Rechnungsdatum", ["date_dmy"]),
        col("/delivery_date", "Lieferdatum", ["date_dmy"]),
        *buyer_fields("Kunde", "Straße", "PLZ", "Ort", "E-Mail"),
        col("/lines/*/description", "Artikel"), col("/lines/*/quantity", "Menge", ["de_decimal"]),
        col("/lines/*/unit", "Einheit", ["unit_map"]),
        col("/lines/*/net_price", "Einzelpreis brutto", ["de_decimal", "gross_to_net"]),
        col("/lines/*/vat_rate", "MwSt %", ["de_decimal"]), col("/lines/*/vat_category", "MwSt %", ["rate_to_category"]),
    ])


# ------------------------------------------------------------ 8 credit-note rows (negative amounts) mixed with invoices
def credit_notes():
    header = ["Nr.", "Datum", "Belegart", "Kunde", "Straße", "PLZ", "Ort", "E-Mail", "Text", "Betrag netto", "Steuersatz",
              "Betrag brutto", "Bezug auf Rechnung"]
    data = [
        ("RE-2026-0101", "02.10.2026", "Rechnung", "baecker", "Wartung Kühlanlage", "255.00", ""),
        ("GS-2026-0007", "12.10.2026", "Gutschrift", "baecker", "Gutschrift Anfahrt (nicht erbracht)", "-45.00", "RE-2026-0101"),
        ("RE-2026-0106", "21.10.2026", "Rechnung", "nordlicht", "Fernwartung Heizungssteuerung", "320.00", ""),
        ("GS-2026-0008", "29.10.2026", "Gutschrift", "mueller", "Preisnachlass Montage", "-150.00", "RE-2026-0102"),
    ]
    rows = [header]
    for num, d, art, b, text, net, ref in data:
        B = BUYERS[b]
        n = Decimal(net)
        rows.append([num, d, art, B["name"], B["street"], B["postcode"], B["city"], B["email"], text, de(n), "19", de(n * Decimal("1.19")), ref])
    (HERE / "credit_notes.csv").write_bytes(csv_text(rows, ";").encode("utf-8"))
    write_expected("credit_notes.csv", "row_per_invoice", "Nr.", [
        col("/number", "Nr."), col("/issue_date", "Datum", ["date_dmy"]),
        *buyer_fields("Kunde", "Straße", "PLZ", "Ort", "E-Mail"),
        col("/preceding_invoice/number", "Bezug auf Rechnung"),
        col("/lines/*/description", "Text"), const("/lines/*/quantity", "1"), const("/lines/*/unit", "C62"),
        col("/lines/*/net_price", "Betrag netto", ["de_decimal", "abs"]),
        col("/lines/*/vat_rate", "Steuersatz", ["de_decimal"]), col("/lines/*/vat_category", "Steuersatz", ["rate_to_category"]),
        col("/source/expected_totals/net", "Betrag netto", ["de_decimal", "abs"]),
        col("/source/expected_totals/gross", "Betrag brutto", ["de_decimal", "abs"]),
    ])


SELLER_PROFILE = {
    "seller": {
        "name": "Ablaufwerk Testfirma Haustechnik GmbH (fiktiv)",
        "street": "Musterweg 1",
        "postcode": "12345",
        "city": "Musterhausen",
        "country": "DE",
        "vat_id": "DE123456789",
        "electronic_address": "rechnung@testfirma-haustechnik.example",
        "electronic_address_scheme": "EM",
        "contact": {"name": "Erika Mustermann", "phone": "+49 30 0000000", "email": "erika.mustermann@testfirma-haustechnik.example"},
    },
    "payment_defaults": {
        "means_code": "58",
        "iban": "DE89370400440532013000",
        "bic": "COBADEFFXXX",
        "account_name": "Ablaufwerk Testfirma Haustechnik GmbH",
        "payment_days": 14,
    },
    "defaults": {"vat_category": "S", "vat_rate": "19", "unit": "C62", "profile": "xrechnung"},
}

FIXTURES = ["lexware_like.csv", "sevdesk_like.csv", "excel_handmade.xlsx", "access_export.csv", "row_per_line.csv",
            "no_vat_column.csv", "mixed_vat.csv", "credit_notes.csv"]


def main() -> None:
    lexware(); sevdesk(); excel(); access(); per_line(); no_vat(); mixed_vat(); credit_notes()
    (HERE / "seller_profile.json").write_text(json.dumps(SELLER_PROFILE, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("wrote", len(FIXTURES), "fixtures + seller_profile.json to", HERE)


if __name__ == "__main__":
    main()
